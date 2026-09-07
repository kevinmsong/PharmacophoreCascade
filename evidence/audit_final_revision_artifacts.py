"""Audit final Word image resolution and inventory current submission artifacts.

This numerical audit deliberately leaves visual approval to page-by-page review.
"""
import hashlib
import io
import json
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SUB = ROOT / 'ACS_Omega_resubmission'
META = ROOT / 'evidence/outputs/alert_disabled_revision'


def digest(path):
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def main():
    clean=Document(SUB/'main.docx')
    highlighted=Document(SUB/'main_highlighted.docx')
    def paragraphs(doc):
        return [' '.join(p.text.split()) for p in doc.paragraphs if p.text.strip()
                and not p.text.startswith('Supporting Information for Review Only.')]
    assert paragraphs(clean)==paragraphs(highlighted), 'Clean and highlighted scientific paragraphs differ'
    def cells(doc):
        return [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
    assert cells(clean)==cells(highlighted), 'Clean and highlighted tables differ'
    images = []
    for name in ['main', 'main_highlighted', 'supporting_information', 'response_to_reviewers']:
        doc = Document(SUB / (name + '.docx'))
        for i, shape in enumerate(doc.inline_shapes, 1):
            rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
            im = Image.open(io.BytesIO(doc.part.related_parts[rid].blob))
            dx, dy = im.width / (shape.width / 914400), im.height / (shape.height / 914400)
            assert min(dx, dy) >= 599, f'{name} image {i} has only {dx:.1f}/{dy:.1f} effective dpi'
            images.append(dict(document=name, image=i, pixels=list(im.size), effective_dpi=[dx, dy]))
        assert (SUB / (name + '.pdf')).exists()
    for path in SUB.glob('*.png'):
        im = Image.open(path)
        assert all(abs(x - 600) < 1 for x in im.info.get('dpi', (0, 0))), str(path)
    assert Image.open(SUB / 'toc_graphic.png').size == (1950, 1050)
    import fitz
    with fitz.open(SUB / 'toc_graphic.pdf') as graphic:
        assert abs(graphic[0].rect.width - 234) < .001 and abs(graphic[0].rect.height - 126) < .001
    (META / 'final_embedded_image_audit.json').write_text(json.dumps(dict(
        status='passed', images=images, clean_highlighted_scientific_text='identical',
        visual_review='separate page-by-page review required'), indent=2))
    roots = [SUB, ROOT / 'results/absolute_floor_1000', ROOT / 'evidence/data/machine_readable',
             ROOT / 'evidence/outputs/absolute_floor', ROOT / 'evidence/outputs/decoy_replicates',
             ROOT / 'evidence/outputs/docking_top10', ROOT / 'evidence/outputs/revision',
             ROOT / 'evidence/outputs/equivalence', ROOT / 'evidence/outputs/ablation_common',
             ROOT / 'evidence/outputs/attrition']
    roots += [ROOT / ('evidence/outputs/benchmark_' + name) for name in
              ['glp1r_full', 'ghsr_full', 'ntsr1_full', 'mdm2_full', 'glp1r_automated', 'glp1r_7ki0']]
    suffixes = {'.csv', '.tsv', '.json', '.npz', '.png', '.pdf', '.svg', '.docx', '.tex', '.md', '.txt', '.yaml', '.sdf', '.pdbqt'}
    paths = set()
    for folder in roots:
        for path in folder.rglob('*'):
            if path.is_file() and path.suffix.lower() in suffixes and not path.name.startswith('_'):
                paths.add(path)
    paths.update(ROOT / name for name in ['README.md', 'ANALYSIS_REPORT.md', 'author_response.md', 'reproduce/README.md'])
    paths.update(ROOT / name / 'SUPERSEDED.md' for name in
                 ['publication', 'ACS_Omega_submission', 'IEEE_Access_submission'])
    for folder in [ROOT/'results',ROOT/'evidence/outputs']:
        paths.update(p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in suffixes)
    raster_rows = []
    for path in sorted(paths):
        if path.suffix.lower() != '.png':
            continue
        with Image.open(path) as im:
            dpi = im.info.get('dpi', (0, 0))
            assert all(abs(x - 600) < 1 for x in dpi), f'Current output figure is not 600 dpi: {path}: {dpi}'
            raster_rows.append(dict(path=str(path.relative_to(ROOT)), pixels=list(im.size), dpi=list(dpi)))
    (META / 'final_output_figure_resolution.json').write_text(json.dumps(dict(
        status='passed', figures=raster_rows), indent=2), encoding='utf-8')
    rows = [dict(path=str(p.relative_to(ROOT)).replace('\\', '/'), bytes=p.stat().st_size,
                 sha256=digest(p)) for p in sorted(paths) if p.exists()]
    (META / 'final_artifact_inventory.json').write_text(json.dumps(dict(
        created_utc=datetime.now(timezone.utc).isoformat(), files=rows,
        scope='Current submission and fresh scientific outputs; historical archives excluded'), indent=2))
    print(f'Passed effective-resolution audit for {len(images)} embedded images; hashed {len(rows)} current artifacts.')


if __name__ == '__main__':
    main()
