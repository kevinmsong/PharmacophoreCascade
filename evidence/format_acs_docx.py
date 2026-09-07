"""Consistent ACS Word typography, table geometry, and figure/caption pagination."""
import io
import re
from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def node(parent,name,attrs):
    old=parent.find(qn(name))
    if old is None:
        old=OxmlElement(name);parent.append(old)
    for key,val in attrs.items(): old.set(qn(key),str(val))
    return old


def format_document(path):
    doc=Document(str(path))
    # Citeproc appends references after all source text. Restore journal order:
    # acknowledgments, references, then a separate graphical-abstract page.
    toc=next((p for p in doc.paragraphs if p.text=='For Table of Contents Use Only'),None)
    refs=[p for p in doc.paragraphs if p.style.name=='Bibliography']
    if toc is not None and refs:
        if not any(p.text=='References' for p in doc.paragraphs):
            toc.insert_paragraph_before('References',style='Heading 1')
        for p in refs: toc._p.addprevious(p._p)
        toc.paragraph_format.page_break_before=True
    for section in doc.sections:
        section.page_width=Inches(8.5);section.page_height=Inches(11)
        section.left_margin=section.right_margin=Inches(1)
        section.top_margin=section.bottom_margin=Inches(1)
        footer=section.footer.paragraphs[0]
        footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if not footer._p.xpath('.//w:fldSimple'):
            field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
            footer._p.append(field)
    for style in doc.styles:
        if style.type==1:
            style.font.name='Times New Roman'
            fonts=style.element.rPr.find(qn('w:rFonts')) if style.element.rPr is not None else None
            if fonts is not None:
                for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                    fonts.attrib.pop(qn('w:'+attr),None)
    for name in ['Normal','Body Text','First Paragraph','Compact']:
        if name in doc.styles:
            st=doc.styles[name];st.font.size=Pt(11)
            st.paragraph_format.line_spacing=1.15
            st.paragraph_format.space_after=Pt(6)
    for name,size in [('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
        st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True
        st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.keep_with_next=True
        st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6)
    for name in ['Caption','Image Caption','Table Caption']:
        if name in doc.styles:
            st=doc.styles[name];st.font.size=Pt(10)
            st.paragraph_format.line_spacing=1.0
            st.paragraph_format.space_after=Pt(7)
    # Keep captions and images together; retain original revision colors.
    for p in doc.paragraphs:
        p.paragraph_format.widow_control=True
        if p.text.startswith('High-throughput native peptide-contact pharmacophore scoring'):
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.space_after=Pt(10)
            for run in p.runs:
                run.font.name='Times New Roman';run.font.size=Pt(15);run.bold=True
        if p.style.name.startswith('Heading'):
            p.paragraph_format.keep_with_next=True
            for run in p.runs: run.font.name='Times New Roman'
        if p._p.xpath('.//w:drawing'):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
        if re.match(r'Table\s+S?\d+[.:]',p.text) and not p._p.xpath('./w:pPr/w:numPr'):
            p.paragraph_format.keep_with_next=True
        if p.style.name=='Bibliography':
            p.paragraph_format.line_spacing=1.0
            p.paragraph_format.space_after=Pt(4)
    for shape in doc.inline_shapes:
        rid=shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        blob=doc.part.related_parts[rid].blob
        im=Image.open(io.BytesIO(blob))
        dpi=im.info.get('dpi',(600,600))[0]
        width=min(6.5,im.width/dpi)
        height=width*im.height/im.width
        if height>7.6: width*=7.6/height;height=7.6
        shape.width=Inches(width);shape.height=Inches(height)
    for table in doc.tables:
        n=len(table.columns)
        table.autofit=False
        weights=[]
        for j in range(n):
            words=[word for row in table.rows for word in row.cells[j].text.split()]
            weights.append(max(6,min(25,max(map(len,words),default=6))))
        if n==2: weights=[17,48]
        widths=[round(9360*w/sum(weights)) for w in weights]
        widths[-1]=9360-sum(widths[:-1])
        pr=table._tbl.tblPr
        node(pr,'w:tblW',{'w:w':9360,'w:type':'dxa'})
        node(pr,'w:tblInd',{'w:w':80,'w:type':'dxa'})
        node(pr,'w:tblLayout',{'w:type':'fixed'})
        margins=node(pr,'w:tblCellMar',{})
        for side in ['top','bottom']: node(margins,'w:'+side,{'w:w':60,'w:type':'dxa'})
        for side in ['left','right']: node(margins,'w:'+side,{'w:w':80,'w:type':'dxa'})
        borders=node(pr,'w:tblBorders',{})
        for side in ['top','bottom']:
            node(borders,'w:'+side,{'w:val':'single','w:sz':6,'w:color':'333333'})
        for side in ['left','right','insideH','insideV']:
            node(borders,'w:'+side,{'w:val':'nil'})
        grid=table._tbl.tblGrid
        for child in list(grid): grid.remove(child)
        for width in widths: node_grid=OxmlElement('w:gridCol');node_grid.set(qn('w:w'),str(width));grid.append(node_grid)
        for i,row in enumerate(table.rows):
            rowpr=row._tr.get_or_add_trPr();node(rowpr,'w:cantSplit',{})
            if i==0: node(rowpr,'w:tblHeader',{})
            for j,cell in enumerate(row.cells):
                cell.width=Inches(widths[j]/1440)
                node(cell._tc.get_or_add_tcPr(),'w:tcW',{'w:w':widths[j],'w:type':'dxa'})
                if i==0:
                    node(cell._tc.get_or_add_tcPr(),'w:shd',{'w:fill':'F0F2F4'})
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1.0
                    p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
                    p.paragraph_format.keep_with_next=(i==0)
                    if i>0 and re.fullmatch(r'[\d.,%+\-/\s()]+',p.text.strip()):
                        p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
                    for run in p.runs:
                        run.font.name='Times New Roman';run.font.size=Pt(9 if n>=6 else 10)
                        if i==0: run.bold=True
    # Preserve image resolution through subsequent Word saves.
    node(doc.settings._element,'w:doNotAutoCompressPictures',{})
    doc.save(str(path))
